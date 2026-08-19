import os
import uuid
from pptx import Presentation

class PPTEngine:
    def __init__(self, output_dir: str = "public/presentations"):
        self.output_dir = output_dir
        if not os.path.exists(self.output_dir):
            os.makedirs(self.output_dir)
            
    def generate_presentation(self, slides_data: list, filename_prefix: str = "presentation") -> str:
        """
        Generates a .pptx file from the provided JSON slide data.
        Returns the relative URL to the generated file (e.g., /presentations/abc-123.pptx)
        """
        prs = Presentation()
        
        # Iterate over slides data
        for i, slide_data in enumerate(slides_data):
            title_text = slide_data.get("title", "Untitled Slide")
            bullets = slide_data.get("bullets", [])
            
            # Use title slide layout for the very first slide (layout 0)
            # Use bullet slide layout for the rest (layout 1)
            layout_index = 0 if i == 0 else 1
            slide_layout = prs.slide_layouts[layout_index]
            slide = prs.slides.add_slide(slide_layout)
            
            # Set Title
            title = slide.shapes.title
            if title:
                title.text = title_text
                
            # Set Bullets
            if layout_index == 1 and bullets:
                # Content placeholder is usually index 1
                try:
                    body_shape = slide.placeholders[1]
                    tf = body_shape.text_frame
                    
                    for j, bullet in enumerate(bullets):
                        if j == 0:
                            tf.text = str(bullet)
                        else:
                            p = tf.add_paragraph()
                            p.text = str(bullet)
                            p.level = 0
                except Exception as e:
                    print(f"Error adding bullets to slide {i}: {e}")
                        
            elif layout_index == 0 and bullets:
                # Subtitle for title slide is usually index 1
                try:
                    subtitle = slide.placeholders[1]
                    if subtitle:
                        subtitle.text = "\n".join(str(b) for b in bullets)
                except Exception as e:
                    print(f"Error adding subtitle to title slide: {e}")
                    
        # Save presentation
        unique_id = uuid.uuid4().hex[:8]
        safe_prefix = "".join(c if c.isalnum() else "_" for c in filename_prefix)
        filename = f"{safe_prefix}_{unique_id}.pptx"
        filepath = os.path.join(self.output_dir, filename)
        
        prs.save(filepath)
        
        # Return URL path
        return f"/presentations/{filename}"
